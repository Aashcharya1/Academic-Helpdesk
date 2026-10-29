"""Explicit workbook field mappings, reviewed by the local administrator."""
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter,column_index_from_string
from ..config import CONFIG

FIELDS={'code','title','department','level','category','credits','ltp','slot','section','instructor','co_instructor','venue','name','email','day','start','end','time','syllabus_ref'}

def map_workbook(path,doc_id,version,spec):
    wb=load_workbook(path,data_only=True)
    output={'records':[],'faculty':[],'slots':[]}
    for sheet,options in spec.items():
        if sheet not in wb:raise ValueError('Unknown worksheet: '+sheet)
        kind=options.get('kind')
        if kind not in ('offering','timetable','faculty','slot'):raise ValueError('Unsupported mapping kind')
        columns=options['columns']
        if not set(columns)<=FIELDS:raise ValueError('Unapproved field name')
        if not all(re.fullmatch('[A-Z]{1,3}',col) for col in columns.values()):raise ValueError('Columns must be Excel letters')
        header=int(options.get('header_row',1));s=wb[sheet]
        if not 0<=header<s.max_row:raise ValueError('Invalid header row')
        for ri in range(header+1,s.max_row+1):
            fields={key:('' if s.cell(ri,column_index_from_string(col)).value is None else str(s.cell(ri,column_index_from_string(col)).value).strip()) for key,col in columns.items()}
            base={'record_id':f'MAP-{doc_id}-{sheet}-{ri}-{kind}','source_doc_id':doc_id,'version':version,'term':CONFIG['term'],'locator':f"'{sheet}'!A{ri}:{get_column_letter(s.max_column)}{ri}"}
            if kind in ('offering','timetable'):
                fields['code']=fields.get('code','').replace(' ','').upper()
                if not re.fullmatch('[A-Z]{2,4}\\d{4,5}',fields['code']):continue
                row={key:'' for key in ('code','title','department','level','category','credits','ltp','slot','section','instructor','co_instructor','venue','syllabus_ref')}
                row.update(fields);row.update(base);row.update(kind=kind,offered_status='listed')
                output['records'].append(row)
            elif kind=='faculty':
                if not fields.get('name') or not re.fullmatch(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',fields.get('email','')):continue
                output['faculty'].append(base|fields)
            else:
                if not fields.get('slot'):continue
                output['slots'].append(base|fields)
    return output
