import json
from pathlib import Path
import pytest
from helpdesk.schemas import Evidence
from helpdesk.service import HelpdeskService,validate_claims
from helpdesk.storage import redact
from helpdesk.tools import draft_email,DataTools
from helpdesk.importers import extract,sha
from helpdesk.config import dataset_dir

@pytest.fixture(autouse=True)
def isolated_test_state(tmp_path,monkeypatch):
    from helpdesk import storage,service
    from filelock import FileLock
    state=tmp_path/'state';state.mkdir()
    monkeypatch.setattr(storage,'dataset_dir',lambda dataset:state)
    monkeypatch.setattr(service,'PROCESS_LOCK',FileLock(str(state/'inference.lock')))

def test_citation_validation():
    e=Evidence(id='E1',doc_id='D1',title='fixture',locator='p.1',text='This is a synthetic source used only for testing.',version='1')
    assert validate_claims({'claims':[{'text':'fixture','evidence_ids':['E1'],'quote':'synthetic source used only for testing'}]},{'E1':e})
    assert not validate_claims({'claims':[{'evidence_ids':['fake'],'quote':'synthetic source'}]},{'E1':e})
    assert not validate_claims({'claims':[{'evidence_ids':['E1'],'quote':'attendance is 50 percent'}]},{'E1':e})
    punctuation=Evidence(id='P1',doc_id='D1',title='fixture',locator='p.1',text='Making prototypes – physical, digital, role-playing.',version='1')
    assert validate_claims({'claims':[{'evidence_ids':['P1'],'quote':'Making prototypes-physical, digital, role - playing.'}]},{'P1':punctuation})
    from helpdesk.service import repair_quote_citations
    misplaced={'claims':[{'evidence_ids':['E1'],'quote':'Making prototypes-physical, digital, role-playing.'}]}
    assert repair_quote_citations(misplaced,{'E1':e,'P1':punctuation})==1
    assert validate_claims(misplaced,{'E1':e,'P1':punctuation})
    fabricated={'claims':[{'evidence_ids':['invented'],'quote':'Making prototypes-physical, digital, role-playing.'}]}
    assert repair_quote_citations(fabricated,{'P1':punctuation})==0

def test_prerequisite_focus_copies_explicit_field_only():
    from helpdesk.service import prerequisite_excerpt,prerequisite_evidence
    assert prerequisite_excerpt('Title Synthetic course Prerequisite None Objective: Synthetic objective.')=='Prerequisite None'
    assert prerequisite_excerpt('Prerequisites: EXM1000 and EXM1001\nLearning Outcomes: synthetic')=='Prerequisites: EXM1000 and EXM1001'
    assert prerequisite_excerpt('Learning Outcomes: Synthetic fixture.')==''
    course=Evidence(id='S1',doc_id='D1',title='Synthetic Course',locator='p.1',text='Number EXL1000 Prerequisite None Objective synthetic.',version='1')
    unrelated=course.model_copy(update={'id':'S2','title':'Other Course','text':'Number EXL2000 Prerequisite None Objective synthetic.'})
    assert prerequisite_evidence([course,unrelated],'What are prerequisites for Synthetic Course?')==[course]
    assert prerequisite_evidence([course,unrelated],'What are prerequisites for EXL1000?')==[course]
    assert prerequisite_evidence([course,unrelated],'What are prerequisites for Missing Course?')==[]

def test_out_of_scope_and_logging():
    r=HelpdeskService().answer('What is the weather tomorrow?',{'dataset':'demo'})
    assert r.status=='out_of_scope'
    from helpdesk.storage import db
    with db('demo') as c: assert c.execute('SELECT COUNT(*) FROM unresolved').fetchone()[0]>0

def test_program_missing():
    r=HelpdeskService().answer('How many degree credits do I need?',{'dataset':'demo','program':''})
    assert r.status=='needs_clarification' and 'program' in r.clarification_fields

def test_dataset_isolation():
    assert dataset_dir('real')!=dataset_dir('demo')
    with pytest.raises(ValueError):dataset_dir('../private_import')

def test_draft_missing_contact():
    d=draft_email('leave','[instructor email]','','')
    assert '[start date]' in d['body'] and '[student name]' in d['body']
    assert 'illness' not in d['body']
    with pytest.raises(ValueError):draft_email('leave','not an email','','')

def test_postponement_draft_preserves_assessment_and_course():
    from helpdesk.tools import draft_academic_request
    draft=draft_academic_request('Write me a draft mail to ask to postpone Quiz 5 of Computer Vision CSL7360','[instructor email]','CSL7360')
    assert 'Quiz 5' in draft['subject'] and 'CSL7360' in draft['subject']
    assert 'postponement of Quiz 5' in draft['body']
    assert 'leave' not in draft['body'] and '[start date]' not in draft['body']
    assert 'illness' not in draft['body']
    other=draft_academic_request('Ask for feedback on my project.','[instructor email]','CSL7360')
    assert 'feedback on my project' in other['body'] and 'leave' not in other['body']

def test_sensitive_log_redaction():
    assert 'diagnosis' not in redact('My diagnosis is private; I need medical leave')
    assert 'someone@iitj.ac.in' not in redact('Who is someone@iitj.ac.in?')

def test_loaders(tmp_path):
    p=tmp_path/'test.csv';p.write_text('code,title\nTEST,Synthetic fixture\n')
    assert list(extract(p))[1]['locator']=='row 2'
    from openpyxl import Workbook
    w=Workbook();w.active['A1']='Synthetic fixture';w.create_sheet('Second')['B2']='Another fixture'
    p=tmp_path/'fixture.xlsx';w.save(p)
    units=list(extract(p));assert any('Second' in u['locator'] and 'B2=' in u['text'] for u in units)
    p=tmp_path/'fixture.eml';p.write_text('Subject: synthetic\nFrom: fixture@example.org\nMessage-ID: <test>\nContent-Type: text/plain\n\nAcademic synthetic fixture')
    assert 'message <test>' in list(extract(p))[0]['locator']

def test_pdf_docx_html_text_loaders_preserve_sources(tmp_path):
    import pymupdf
    from docx import Document
    pdf=tmp_path/'fixture.pdf';document=pymupdf.open();page=document.new_page();page.insert_text((72,72),'Synthetic academic fixture, not an institutional rule.');document.save(pdf);document.close()
    docx=tmp_path/'fixture.docx';document=Document();document.add_paragraph('Synthetic academic fixture.');document.save(docx)
    html=tmp_path/'fixture.html';html.write_text('<html><body><h1>Synthetic academic fixture</h1></body></html>')
    text=tmp_path/'fixture.txt';text.write_text('Synthetic academic fixture.')
    for source in (pdf,docx,html,text):
        original=sha(source);units=list(extract(source))
        assert units and any('Synthetic academic fixture' in unit['text'] for unit in units)
        assert all(unit['locator'] for unit in units) and sha(source)==original

def test_duplicate_hash(tmp_path):
    p=tmp_path/'a';q=tmp_path/'b';p.write_bytes(b'fixture');q.write_bytes(b'fixture')
    assert sha(p)==sha(q)

def test_budget_limits():
    from helpdesk.agents import Budget
    b=Budget()
    for _ in range(6):b.check('data')
    with pytest.raises(RuntimeError):b.check('data')

def test_exact_tools():
    tool=object.__new__(DataTools)
    tool.context={'term':'fixture'};tool.review={};tool.faculty=[];tool.slots=[]
    tool.records=[{'code':'TEST1000','title':'Synthetic fixture','term':'fixture','kind':'offering','slot':'A','venue':'','instructor':'Fixture Person','co_instructor':'','department':'Test','level':'','category':''}]
    tool.evidence=lambda row:None
    assert tool.schedule('TEST1000').status=='ok'
    assert tool.schedule('Unknown').status=='not_found'
    assert tool.instructor('TEST1000').data['recipient']=='[instructor email]'
    assert tool.count('Test','100','elective').status=='incomplete'
    tool.records.append({**tool.records[0],'slot':'B'})
    assert tool.schedule('TEST1000').status=='ambiguous'

def test_no_source_refusal(monkeypatch):
    from helpdesk import service
    class EmptyCorpus:
        def __init__(self,*args):self.chunks=[]
    monkeypatch.setattr(service,'Retriever',EmptyCorpus)
    assert HelpdeskService().answer('What are the academic registration rules?',{'dataset':'demo'}).status=='insufficient_evidence'

def test_manual_mapping_retains_zero_and_provenance(tmp_path):
    from openpyxl import Workbook
    from helpdesk.importers.mapping import map_workbook
    w=Workbook();w.active.title='Fixture';w.active.append(['Code','Name','Credit']);w.active.append(['EXM1000','Synthetic',0])
    p=tmp_path/'fixture.xlsx';w.save(p)
    rows=map_workbook(p,'DEMO','fixture',{'Fixture':{'kind':'offering','header_row':1,'columns':{'code':'A','title':'B','credits':'C'}}})['records']
    assert rows[0]['credits']=='0' and rows[0]['locator']=="'Fixture'!A2:C2"

def test_eml_attachment_deduplication(tmp_path,monkeypatch):
    from email.message import EmailMessage
    from helpdesk import importers
    ingest=importers.ingest
    corpus=tmp_path/'isolated-demo';corpus.mkdir()
    monkeypatch.setattr(importers,'dataset_dir',lambda dataset:corpus)
    m=EmailMessage();m['Subject']='Synthetic fixture';m['Message-ID']='<synthetic-fixture>';m.set_content('Synthetic academic fixture')
    m.add_attachment(b'Synthetic attachment',maintype='text',subtype='plain',filename='attachment.txt')
    (tmp_path/'a.eml').write_bytes(m.as_bytes());(tmp_path/'duplicate.txt').write_bytes(b'Synthetic attachment')
    result=ingest('demo',tmp_path,use_ocr=False)
    assert result['files']==3 and result['distinct']==2

def test_session_context_boundary():
    from helpdesk.storage import save_session,load_session
    save_session('demo','fixture-session','snapshot-A','EXM1000','course schedule',['DEMO-E1'])
    assert load_session('demo','fixture-session','snapshot-A')['course']=='EXM1000'
    assert load_session('demo','fixture-session','snapshot-B')=={}

def test_exact_count_deduplicates_verified_offerings():
    tool=object.__new__(DataTools);tool.context={'term':'fixture'}
    tool.review={'complete_departments':['Synthetic department']}
    row={'kind':'offering','term':'fixture','department':'Synthetic department','level':'400','category':'elective','code':'EXM4000'}
    tool.records=[row,row.copy(),dict(row,code='EXM4001')]
    tool.evidence=lambda row:None
    assert tool.count('Synthetic department','400','elective').data['count']==2
    tool.review={'complete_departments':[]}
    assert tool.count('Synthetic department','400','elective').status=='incomplete'

def test_edit_candidate_preserves_active_snapshot(tmp_path,monkeypatch):
    from helpdesk import importers
    root=tmp_path/'demo';active=root/'snapshots'/'fixture';active.mkdir(parents=True)
    (active/'review.json').write_text('{"approved":[]}')
    (active/'chroma').mkdir();(active/'chroma'/'database').write_text('active index')
    (root/'active.txt').write_text(str(active));(root/'candidate.txt').write_text(str(active))
    monkeypatch.setattr(importers,'dataset_dir',lambda dataset:root)
    candidate=importers.editable_candidate('demo')
    assert candidate!=active and not (candidate/'chroma').exists()
    (candidate/'review.json').write_text('{"approved":["synthetic"]}')
    assert (active/'review.json').read_text()=='{"approved":[]}'

@pytest.mark.parametrize('malformed', [False,True])
def test_draft_survives_model_status_and_json_errors(tmp_path,monkeypatch,malformed):
    from helpdesk import service,agents
    from helpdesk.schemas import ToolResult
    (tmp_path/'index.json').write_text('{}')
    class Corpus:
        def __init__(self,*args):self.path=tmp_path;self.chunks=[{'text':'fixture'}]
    class ExactTools:
        def __init__(self,*args):self.records=[]
        def instructor(self,*args):return ToolResult(status='not_found')
    class Api:
        def raise_for_status(self):pass
        def json(self):return {'models':[]}
    async def execute(question,context,mode,model,tools,budget,*args):
        tools['create_email_draft']('sick leave','EXM1000','invented diagnosis','')
        return '{"status":' if malformed else '{"status":"answered","claims":[]}'
    monkeypatch.setattr(service,'Retriever',Corpus)
    monkeypatch.setattr(service,'DataTools',ExactTools)
    monkeypatch.setattr(agents,'execute',execute)
    monkeypatch.setattr('requests.get',lambda *a,**kw:Api())
    result=HelpdeskService().answer('Draft leave for EXM1000 for 2 days.',{'dataset':'demo'},'single')
    assert result.status=='draft' and result.draft['recipient']=='[instructor email]'
    assert 'invented diagnosis' not in result.draft['body']
    assert 'sick' not in result.draft['body']

@pytest.mark.parametrize('tool_status,expected',[('ok','insufficient_evidence'),('ambiguous','needs_clarification')])
def test_unprovided_model_section_does_not_hide_records(tmp_path,monkeypatch,tool_status,expected):
    from helpdesk import service,agents
    from helpdesk.schemas import ToolResult
    (tmp_path/'index.json').write_text('{}')
    class Corpus:
        def __init__(self,*args):self.path=tmp_path;self.chunks=[{'text':'fixture'}]
    class ExactTools:
        def __init__(self,*args):self.records=[]
        def schedule(self,course,section):
            assert section==''
            return ToolResult(status=tool_status,data={'courses':[],'meetings':[]})
    class Api:
        def raise_for_status(self):pass
        def json(self):return {'models':[]}
    async def execute(question,context,mode,model,tools,budget,*args):
        tools['lookup_course_schedule']('Synthetic subject','invented section')
        return '{"result":""}'
    monkeypatch.setattr(service,'Retriever',Corpus)
    monkeypatch.setattr(service,'DataTools',ExactTools)
    monkeypatch.setattr(agents,'execute',execute)
    monkeypatch.setattr('requests.get',lambda *a,**kw:Api())
    result=HelpdeskService().answer('What is the schedule for Synthetic subject?',{'dataset':'demo'},'single')
    assert result.status==expected and 'error_type' not in result.metrics

@pytest.mark.parametrize('model_k,expected_k',[('3',3),('20',4),('-2',1),('invalid',4),(None,4)])
def test_model_search_count_is_normalized(tmp_path,monkeypatch,model_k,expected_k):
    from helpdesk import service,agents
    searches=[]
    (tmp_path/'index.json').write_text('{}')
    class Corpus:
        def __init__(self,*args):self.path=tmp_path;self.chunks=[{'text':'fixture'}]
        def search(self,question,k):
            assert k==expected_k
            assert question.startswith('Find synthetic course contents.')
            searches.append(question)
            return []
    class ExactTools:
        def __init__(self,*args):self.records=[]
    class Api:
        def raise_for_status(self):pass
        def json(self):return {'models':[]}
    async def execute(question,context,mode,model,tools,budget,*args):
        tools['search_documents']('Synthetic academic fixture',model_k)
        if model_k=='3':
            tools['search_documents']('A reformulation',model_k)
            assert tools['search_documents']('A third attempt',model_k)['status']=='incomplete'
        return '{"status":"insufficient_evidence","claims":[]}'
    monkeypatch.setattr(service,'Retriever',Corpus)
    monkeypatch.setattr(service,'DataTools',ExactTools)
    monkeypatch.setattr(agents,'execute',execute)
    monkeypatch.setattr('requests.get',lambda *a,**kw:Api())
    result=HelpdeskService().answer('Find synthetic course contents.',{'dataset':'demo'},'single')
    assert result.status=='insufficient_evidence' and 'error_type' not in result.metrics
    assert len(searches)==(2 if model_k=='3' else 1)

def test_failed_smoke_recheck_keeps_actual_prior_results(tmp_path,monkeypatch):
    from helpdesk import smoke as checks
    from types import SimpleNamespace
    calls=[];fail_policy=True
    class SyntheticService:
        def answer(self,question,*args,**kwargs):
            calls.append(question)
            is_policy='prototyping' in question
            tools=['search_documents'] if is_policy else ['lookup_course_schedule','create_email_draft'] if 'draft' in question else ['create_email_draft'] if 'email' in question else ['lookup_course_schedule']
            row={'status':'error' if is_policy and fail_policy else 'draft' if 'create_email_draft' in tools else 'answered','citations':[{'title':'Design Thinking synthetic test fixture'}] if is_policy else [],'metrics':{'tool_trace':[{'tool':t} for t in tools]}}
            return SimpleNamespace(model_dump=lambda:row)
    monkeypatch.setattr(checks,'ROOT',tmp_path)
    monkeypatch.setattr(checks,'HelpdeskService',SyntheticService)
    (tmp_path/'.cache').mkdir()
    assert not checks.smoke('fixture','single')['passed']
    calls.clear();fail_policy=False
    result=checks.smoke('fixture','single',failed_only=True)
    assert result['passed'] and result['executed_now']==['policy'] and len(calls)==1
    assert set(result['reused_checks'])=={'schedule','email','compound'}
    saved=json.loads((tmp_path/'report/private/smoke-fixture-single.json').read_text())
    assert all(row['passed'] for row in saved)
    assert all(row.get('reused_check',{}).get('artifact_sha256') for row in saved if row['case']!='policy')
    calls.clear()
    result=checks.smoke('fixture','single',recheck_cases=['policy'])
    assert result['passed'] and result['executed_now']==['policy'] and len(calls)==1

def test_post_retrieval_request_disables_tools_and_requires_quotes():
    import asyncio
    from google.genai import types
    from google.adk.models.llm_request import LlmRequest
    from google.adk.models.lite_llm import _get_completion_inputs
    from helpdesk.agents import Budget
    config=types.GenerateContentConfig(tools=[types.Tool(function_declarations=[types.FunctionDeclaration(name='search_documents',description='synthetic fixture')])])
    request=LlmRequest(model='ollama_chat/llama3.1:8b',contents=[types.Content(role='tool',parts=[types.Part.from_function_response(name='search_documents',response={'status':'ok','evidence':[{'id':'SYNTHETIC','text':'Synthetic fixture.'}]})])],config=config)
    async def check():
        await Budget().before_model(None,request)
        _,tools,format,_,_=await _get_completion_inputs(request,model='ollama_chat/llama3.1:8b')
        assert not tools
        assert format['type']=='json_schema'
        assert set(format['json_schema']['schema']['properties']['claims']['items']['required'])=={'evidence_ids','quote'}
        assert config.tools # Per-generation changes must not leak to the shared agent config.
    asyncio.run(check())
